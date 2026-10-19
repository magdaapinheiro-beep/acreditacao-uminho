# -*- coding: utf-8 -*-
"""
importar_tabela.py — Importação automática a partir de Excel (.xlsx) ou CSV
===========================================================================

Lê os ficheiros de tabela depositados na pasta fontes/ e atualiza o dados.csv
(ciclos de estudos) e o institucional.csv (acreditação institucional). É a
alternativa ao PDF da A3ES: em vez de depositar deliberações, deposita-se o
levantamento em Excel ou CSV. É chamado pelo atualizar.py, mas também corre à
parte:

    python3 importar_tabela.py

Como funciona
-------------
- Aceita .xlsx (todas as folhas) e .csv (separador ; ou ,).
- Reconhece as colunas pelos cabeçalhos, tanto no formato do levantamento da
  UMinho ("Designação do Ciclo de estudos", "N.º de processo A3ES", ...) como
  nos nomes internos do dados.csv (nome, proc, dataDecisao, ...).
- No .xlsx, os links (A3ES / IES) são lidos do hyperlink da célula, não do texto.
- Distingue automaticamente a tabela de ciclos (tem "processo"/"designação") da
  tabela institucional (tem "instituição" e não tem designação de ciclo).

Regra de fusão (a TABELA manda)
-------------------------------
Quando um processo já existe no dados.csv, os campos preenchidos na tabela
substituem os que lá estão (a tabela é a fonte oficial); campos vazios na tabela
NÃO apagam o que já existe. As linhas importadas ficam com origem=excel|csv e
validado=sim.

A tabela também manda no que SAI: um ciclo vindo de Excel/CSV que já não conste
de nenhuma tabela em fontes/ é retirado do dados.csv (ex.: uma UO que não
autorizou a publicação). Fica cópia em dados.csv.bak-antes-retirada-AAAAMMDD.
As linhas vindas de PDF (origem=pdf) nunca são retiradas por esta regra.

Correções registadas (correcoes.csv)
------------------------------------
Erros conhecidos da tabela de origem que ainda não foram corrigidos lá. Cada
linha (proc;campo;valor;motivo) substitui esse campo do ciclo ANTES da fusão,
pelo que sobrevive a todas as reimportações. Quando a tabela de origem já trouxer
o valor certo, o relatório avisa que a correção pode ser apagada.
Links A3ES/DGES trocados entre colunas são corrigidos sozinhos (e assinalados).
"""

import io, os, re, csv, sys, datetime, unicodedata

AQUI   = os.path.dirname(os.path.abspath(__file__))
FONTES = os.path.join(AQUI, "fontes")
DELIB  = os.path.join(AQUI, "deliberacoes")   # tolerância: tabelas largadas aqui também são lidas
DADOS  = os.path.join(AQUI, "dados.csv")
INST   = os.path.join(AQUI, "institucional.csv")
CORR   = os.path.join(AQUI, "correcoes.csv")

# Unidades orgânicas da UMinho (código -> nome oficial), para normalizar a coluna uo
UOS = [
    ("EAAD", "Escola de Arquitetura, Arte e Design"),
    ("EC",   "Escola de Ciências"),
    ("ED",   "Escola de Direito"),
    ("EE",   "Escola de Engenharia"),
    ("EEG",  "Escola de Economia e Gestão"),
    ("ELACH","Escola de Letras, Artes e Ciências Humanas"),
    ("EM",   "Escola de Medicina"),
    ("EP",   "Escola de Psicologia"),
    ("ESE",  "Escola Superior de Enfermagem"),
    ("ICS",  "Instituto de Ciências Sociais"),
    ("IE",   "Instituto de Educação"),
    ("I3Bs", "Instituto de Investigação em Biomateriais, Biodegradáveis e Biomiméticos"),
]

COLUNAS_CICLO = ["nome","grau","uo","estado","proc","dataDecisao","anos","fim",
                 "registo","dataRegisto","a3es","dges","origem","validado"]
COLUNAS_INST  = ["entidade","tipo","estado","nivel","validade","ref","linkA3es","linkDges"]

# Cabeçalho normalizado (ver _norm) -> campo interno
ALIASES = {
    "designacao do ciclo de estudos": "nome", "designacao do ciclo": "nome", "nome": "nome",
    "grau": "grau",
    "unidade organica": "uo", "uo": "uo",
    "estado": "estado",
    "n o de processo a3es": "proc", "numero de processo a3es": "proc",
    "processo a3es": "proc", "processo": "proc", "proc": "proc",
    "data de decisao de acreditacao": "dataDecisao", "data de decisao": "dataDecisao",
    "datadecisao": "dataDecisao",
    "anos": "anos",
    "termino de acreditacao": "fim", "termino": "fim", "fim": "fim",
    "registo no ies": "registo", "registo": "registo",
    "data do registo": "dataRegisto", "dataregisto": "dataRegisto",
    "link a3es": "a3es", "a3es": "a3es",
    "link ies": "dges", "link dges": "dges", "linkdges": "dges", "dges": "dges",
    "link dges ies": "dges", "link ies dges": "dges",
    # institucional
    "instituicao": "entidade", "entidade": "entidade",
    "data de acreditacao": "inicio",
    "validade": "validade", "nivel": "nivel", "ref": "ref",
}


def _txt(v):
    if v is None:
        return ""
    return re.sub(r"\s+", " ", str(v).replace("\xa0", " ")).strip()


def _norm(s):
    s = unicodedata.normalize("NFKD", _txt(s))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", " ", s.lower()).strip()


def _data(v):
    """datetime, '22/04/2025' ou '2025-04-22' -> AAAA-MM-DD."""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    s = _txt(v)
    if not s:
        return ""
    m = re.search(r"(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        return m.group(0)
    m = re.search(r"(\d{1,2})/(\d{1,2})/(\d{4})", s)
    if m:
        return "%s-%02d-%02d" % (m.group(3), int(m.group(2)), int(m.group(1)))
    return s


def _key(proc):
    return re.sub(r"[^A-Za-z0-9]", "", proc or "").upper()


def _link(get_link, get, campo):
    """URL do hyperlink; senão o texto da célula, mas só se for mesmo um URL.
    Descarta placeholders como 'Consulta do processo' / 'Ver Registo'."""
    u = get_link(campo)
    if u:
        return u
    t = get(campo)
    return t if re.match(r"^https?://", t or "") else ""


def _uo(valor):
    s = _txt(valor)
    if not s:
        return ""
    base = re.sub(r"\s*\([^)]*\)\s*$", "", s)     # remove sufixo "(ED)"
    alvo = _norm(base)
    for cod, nome in UOS:
        n = _norm(nome)
        if alvo and (n in alvo or alvo in n):
            return "%s — %s" % (cod, nome)
    m = re.search(r"\(([A-Za-z0-9]+)\)", s)        # tenta o código entre parênteses
    if m:
        code = m.group(1).upper()
        for cod, nome in UOS:
            if cod.upper() == code:
                return "%s — %s" % (cod, nome)
    return base


# ---------------------------------------------------------------------------
# Leitura das fontes: cada tabela vira {"ciclos": [...], "inst": [...]}
# ---------------------------------------------------------------------------

def _classificar(campos):
    if "nome" in campos or ("proc" in campos and "entidade" not in campos):
        return "ciclos"
    if "entidade" in campos:
        return "inst"
    if "proc" in campos:
        return "ciclos"
    return None


def _montar_ciclo(get, get_link, origem):
    rec = {
        "nome":        get("nome"),
        "grau":        get("grau"),
        "uo":          _uo(get("uo")),
        "estado":      get("estado"),
        "proc":        get("proc"),
        "dataDecisao": _data(get("dataDecisao", bruto=True)),
        "anos":        (re.sub(r"\D", "", get("anos")) or get("anos")),
        "fim":         _data(get("fim", bruto=True)),
        "registo":     get("registo"),
        "dataRegisto": _data(get("dataRegisto", bruto=True)),
        "a3es":        _link(get_link, get, "a3es"),
        "dges":        _link(get_link, get, "dges"),
        "origem":      origem,
        "validado":    "sim",
    }
    return rec if rec["proc"] else None


def _montar_inst(get, get_link, origem):
    ini = _data(get("inicio", bruto=True))
    fim = _data(get("fim", bruto=True))
    validade = "%s-%s" % (ini[:4], fim[:4]) if (ini[:4] and fim[:4]) else ""
    rec = {
        "entidade": get("entidade"),
        "tipo":     "Acreditação institucional (A3ES)",
        "estado":   get("estado") or "Acreditada",
        "nivel":    get("nivel") or "ok",
        "validade": get("validade") or validade,
        "ref":      get("ref") or get("proc"),
        "linkA3es": _link(get_link, get, "a3es"),
        "linkDges": _link(get_link, get, "dges"),
    }
    return rec if rec["entidade"] else None


def _processar_tabela(cabecalho, linhas_val, linhas_link, origem):
    """cabecalho: list[str]; linhas_val: list[list]; linhas_link: list[dict campo->url]."""
    idx = {}
    for j, h in enumerate(cabecalho):
        campo = ALIASES.get(_norm(h))
        if campo and campo not in idx:
            idx[campo] = j
    tipo = _classificar(set(idx))
    if not tipo:
        return {"ciclos": [], "inst": []}

    out = {"ciclos": [], "inst": []}
    for i, linha in enumerate(linhas_val):
        def get(campo, bruto=False):
            j = idx.get(campo)
            if j is None or j >= len(linha):
                return ""
            v = linha[j]
            return v if bruto else _txt(v)
        def get_link(campo):
            return (linhas_link[i].get(campo) if i < len(linhas_link) else "") or ""
        rec = (_montar_ciclo if tipo == "ciclos" else _montar_inst)(get, get_link, origem)
        if rec:
            out[tipo].append(rec)
    return out


def _ler_xlsx(caminho):
    try:
        import openpyxl
    except ImportError:
        print("  ! %s ignorado (falta openpyxl: pip install openpyxl)" % os.path.basename(caminho))
        return {"ciclos": [], "inst": []}
    import warnings
    warnings.filterwarnings("ignore")
    origem = "excel"
    acc = {"ciclos": [], "inst": []}
    wb = openpyxl.load_workbook(caminho, data_only=True)
    for ws in wb.worksheets:
        filas = list(ws.iter_rows())
        if not filas:
            continue
        cabecalho = [_txt(c.value) for c in filas[0]]
        linhas_val, linhas_link = [], []
        for fila in filas[1:]:
            if not any(_txt(c.value) for c in fila):
                continue
            linhas_val.append([c.value for c in fila])
            link = {}
            for j, c in enumerate(fila):
                if c.hyperlink and c.hyperlink.target:
                    campo = ALIASES.get(_norm(cabecalho[j])) if j < len(cabecalho) else None
                    if campo in ("a3es", "dges"):
                        link[campo] = c.hyperlink.target
            linhas_link.append(link)
        res = _processar_tabela(cabecalho, linhas_val, linhas_link, origem)
        acc["ciclos"].extend(res["ciclos"])
        acc["inst"].extend(res["inst"])
    return acc


def _ler_csv(caminho):
    with io.open(caminho, "r", encoding="utf-8-sig", newline="") as f:
        amostra = f.read(4096)
        f.seek(0)
        delim = ";" if amostra.count(";") >= amostra.count(",") else ","
        filas = list(csv.reader(f, delimiter=delim))
    if not filas:
        return {"ciclos": [], "inst": []}
    cabecalho = filas[0]
    linhas_val = [fila for fila in filas[1:] if any(_txt(x) for x in fila)]
    linhas_link = [{} for _ in linhas_val]
    return _processar_tabela(cabecalho, linhas_val, linhas_link, "csv")


# ---------------------------------------------------------------------------
# Escrita / fusão
# ---------------------------------------------------------------------------

def _ler_dados(caminho, colunas):
    linhas, campos = [], list(colunas)
    if os.path.exists(caminho):
        with io.open(caminho, "r", encoding="utf-8-sig", newline="") as f:
            r = csv.DictReader(f, delimiter=";")
            campos = list(dict.fromkeys(list(r.fieldnames or []) + colunas))
            linhas = [dict(x) for x in r]
    return linhas, campos


def _gravar(caminho, linhas, campos):
    with io.open(caminho, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=campos, delimiter=";", extrasaction="ignore")
        w.writeheader()
        for l in linhas:
            w.writerow({k: l.get(k, "") for k in campos})


def _fundir(caminho, colunas, novos, chave):
    """A tabela manda: sobrepõe campos preenchidos; não apaga com vazios."""
    linhas, campos = _ler_dados(caminho, colunas)
    for c in colunas:
        if c not in campos:
            campos.append(c)
    por_chave = {chave(l): l for l in linhas}
    nnovos = natual = 0
    for rec in novos:
        k = chave(rec)
        if not k:
            continue
        if k in por_chave:
            l = por_chave[k]
            mudou = False
            for campo, val in rec.items():
                if val and (l.get(campo) or "") != val:
                    l[campo] = val; mudou = True
            if mudou:
                natual += 1
        else:
            nova = {c: "" for c in campos}
            nova.update(rec)
            linhas.append(nova)
            por_chave[k] = nova
            nnovos += 1
    _gravar(caminho, linhas, campos)
    return nnovos, natual


def _chave_ciclo(ciclos):
    """Chave de fusão: o n.º de processo A3ES.

    Se a tabela de origem repetir um n.º de processo, esse número deixa de
    distinguir os ciclos e um sobreporia o outro. Nesse caso — e só nesse —
    junta-se o nome e o grau, para que nenhum ciclo se perca. O erro continua
    a ser assinalado por _validar_ciclos: quem o corrige é a tabela.
    """
    repetidos = set()
    vistos = set()
    for r in ciclos:
        k = _key(r.get("proc"))
        if not k:
            continue
        if k in vistos:
            repetidos.add(k)
        vistos.add(k)

    def chave(r):
        k = _key(r.get("proc"))
        if k and k in repetidos:
            return "%s|%s|%s" % (k, _norm(r.get("nome")), _norm(r.get("grau")))
        return k

    return chave


def _validar_ciclos(ciclos):
    """Problemas na tabela de origem que fariam perder ou empobrecer ciclos.

    A fusão usa o n.º de processo A3ES como chave: se dois ciclos diferentes
    partilharem o mesmo número, um sobrepõe o outro e desaparece em silêncio.
    """
    avisos = []

    por_proc = {}
    sem_proc = []
    for r in ciclos:
        k = _key(r.get("proc"))
        rotulo = "%s (%s)" % (r.get("nome") or "?", r.get("grau") or "?")
        if not k:
            sem_proc.append(rotulo)
        else:
            por_proc.setdefault(k, []).append(rotulo)

    colisoes = {k: v for k, v in por_proc.items() if len(v) > 1}
    if colisoes:
        avisos.append("")
        avisos.append("  ATENÇÃO: n.º de processo A3ES repetido em %d caso(s)." % len(colisoes))
        avisos.append("  Os ciclos foram todos importados, mas o número (e o link A3ES)")
        avisos.append("  está errado num dos de cada par. Corrigir na tabela de origem:")
        for k, v in sorted(colisoes.items()):
            avisos.append("    %s" % k)
            for rotulo in v:
                avisos.append("      - %s" % rotulo)

    if sem_proc:
        avisos.append("")
        avisos.append("  ATENÇÃO: %d ciclo(s) sem n.º de processo A3ES (não são importados):" % len(sem_proc))
        for rotulo in sem_proc:
            avisos.append("      - %s" % rotulo)

    sem_link = [("%s (%s)" % (r.get("nome") or "?", r.get("grau") or "?"),
                 [rot for campo, rot in (("a3es", "A3ES"), ("dges", "IES"))
                  if not (r.get(campo) or "").strip()])
                for r in ciclos]
    sem_link = [(nome, falta) for nome, falta in sem_link if falta]
    if sem_link:
        avisos.append("")
        avisos.append("  Aviso: %d ciclo(s) sem link (célula sem hiperligação na tabela):" % len(sem_link))
        for nome, falta in sem_link:
            avisos.append("      - %-45s sem link %s" % (nome[:45], " e ".join(falta)))

    return avisos


def _limpar_duplicados(caminho, colunas):
    """Tira do dados.csv linhas repetidas e assinala processos repetidos.

    Quando a tabela de origem corrige um n.º de processo que antes estava
    errado, a fusão sobrepõe a linha antiga e fica uma cópia exacta da linha
    certa — o ciclo aparecia duas vezes na página. Remove-se aqui a cópia.
    """
    linhas, campos = _ler_dados(caminho, colunas)
    vistas, limpas, removidas = set(), [], []
    for l in linhas:
        assinatura = tuple((c, (l.get(c) or "").strip()) for c in campos)
        if assinatura in vistas:
            removidas.append("%s (%s)" % (l.get("nome") or "?", l.get("grau") or "?"))
            continue
        vistas.add(assinatura)
        limpas.append(l)

    por_proc = {}
    for l in limpas:
        k = _key(l.get("proc"))
        if k:
            por_proc.setdefault(k, []).append(
                "%s (%s)" % (l.get("nome") or "?", l.get("grau") or "?"))
    repetidos = {k: v for k, v in por_proc.items() if len(v) > 1}

    if removidas:
        _gravar(caminho, limpas, campos)

    avisos = []
    if removidas:
        avisos.append("")
        avisos.append("  Removida(s) %d linha(s) repetida(s) do dados.csv:" % len(removidas))
        for rotulo in removidas:
            avisos.append("      - %s" % rotulo)
    if repetidos:
        avisos.append("")
        avisos.append("  ATENÇÃO: %d n.º de processo repetido(s) no dados.csv "
                      "(ciclos diferentes com o mesmo número):" % len(repetidos))
        for k, v in sorted(repetidos.items()):
            avisos.append("    %s" % k)
            for rotulo in v:
                avisos.append("      - %s" % rotulo)
    return avisos


def _destrocar_links(ciclos):
    """Link A3ES e link DGES trocados entre colunas -> repõe-nos no sítio."""
    avisos = []
    for r in ciclos:
        a, d = r.get("a3es") or "", r.get("dges") or ""
        if "dges.gov.pt" in a and "a3es.pt" in d:
            r["a3es"], r["dges"] = d, a
            avisos.append("      - %s (%s)" % (r.get("nome") or "?", r.get("grau") or "?"))
    if avisos:
        avisos = ["", "  Corrigido: %d ciclo(s) com os links A3ES e DGES trocados na tabela "
                      "(corrigir também na origem):" % len(avisos)] + avisos
    return avisos


def _aplicar_correcoes(ciclos):
    """Aplica o correcoes.csv aos ciclos lidos da tabela, antes da fusão."""
    if not os.path.exists(CORR):
        return []
    with io.open(CORR, "r", encoding="utf-8-sig", newline="") as f:
        correcoes = [dict(x) for x in csv.DictReader(f, delimiter=";")]
    por_proc = {}
    for r in ciclos:
        por_proc.setdefault(_key(r.get("proc")), []).append(r)
    aplicadas, redundantes, orfas = [], [], []
    for c in correcoes:
        k, campo, valor = _key(c.get("proc")), _txt(c.get("campo")), _txt(c.get("valor"))
        if not k or not campo:
            continue
        alvos = por_proc.get(k)
        if not alvos:
            orfas.append(c.get("proc"))
            continue
        for r in alvos:
            if (r.get(campo) or "") == valor:
                redundantes.append("%s %s" % (c.get("proc"), campo))
            else:
                aplicadas.append("%s (%s): %s %s -> %s" % (r.get("nome"), r.get("grau"), campo,
                                                          r.get(campo) or "(vazio)", valor))
                r[campo] = valor
    avisos = []
    if aplicadas:
        avisos += ["", "  Correções registadas aplicadas (correcoes.csv):"]
        avisos += ["      - %s" % a for a in aplicadas]
    if redundantes:
        avisos += ["", "  A tabela já traz o valor certo; pode apagar do correcoes.csv:"]
        avisos += ["      - %s" % a for a in redundantes]
    if orfas:
        avisos += ["", "  Correções sem ciclo correspondente na tabela (verificar o n.º de processo):"]
        avisos += ["      - %s" % a for a in orfas]
    return avisos


def _validar_datas(ciclos):
    """Fim da acreditação anterior à própria decisão: quase sempre erro de cópia."""
    maus = ["      - %s (%s) %s: decisão %s, fim %s" % (r.get("nome"), r.get("grau"), r.get("proc"),
                                                    r.get("dataDecisao"), r.get("fim"))
            for r in ciclos
            if r.get("fim") and r.get("dataDecisao") and r["fim"] < r["dataDecisao"]]
    if maus:
        maus = ["", "  ATENÇÃO: %d ciclo(s) com fim da acreditação anterior à decisão "
                    "(corrigir na tabela ou no correcoes.csv):" % len(maus)] + maus
    return maus


def _retirar_ausentes(caminho, colunas, ciclos, chave):
    """Retira do dados.csv os ciclos de Excel/CSV que já não estão nas tabelas."""
    na_fonte = {chave(r) for r in ciclos}
    linhas, campos = _ler_dados(caminho, colunas)
    ficam, saem = [], []
    for l in linhas:
        if (l.get("origem") or "") in ("excel", "csv") and chave(l) not in na_fonte:
            saem.append(l)
        else:
            ficam.append(l)
    if not saem:
        return []
    copia = "%s.bak-antes-retirada-%s" % (caminho, datetime.date.today().strftime("%Y%m%d"))
    if not os.path.exists(copia):
        import shutil
        shutil.copy2(caminho, copia)
    _gravar(caminho, ficam, campos)
    avisos = ["", "  Retirado(s) %d ciclo(s) que já não constam da tabela de origem "
                  "(cópia em %s):" % (len(saem), os.path.basename(copia))]
    por_uo = {}
    for l in saem:
        por_uo.setdefault(l.get("uo") or "?", []).append(l)
    for uo, ls in sorted(por_uo.items()):
        avisos.append("      %s — %d ciclo(s)" % (uo.split(" — ")[0], len(ls)))
    return avisos


def _localizar_fontes():
    """Procura .xlsx/.csv em fontes/ e, por tolerância, em deliberacoes/."""
    achados = []
    for pasta in (FONTES, DELIB):
        if not os.path.isdir(pasta):
            continue
        for f in sorted(os.listdir(pasta)):
            if f.lower().endswith((".xlsx", ".csv")):
                achados.append(os.path.join(pasta, f))
    return achados


def importar(verbose=True):
    caminhos = _localizar_fontes()
    if not caminhos:
        return 0

    ciclos, inst = [], []
    for caminho in caminhos:
        res = (_ler_xlsx(caminho) if caminho.lower().endswith(".xlsx")
               else _ler_csv(caminho))
        ciclos.extend(res["ciclos"])
        inst.extend(res["inst"])

    avisos = _validar_ciclos(ciclos)
    avisos += _destrocar_links(ciclos)
    avisos += _aplicar_correcoes(ciclos)
    avisos += _validar_datas(ciclos)

    chave = _chave_ciclo(ciclos)
    nc, ac = _fundir(DADOS, COLUNAS_CICLO, ciclos, chave)
    ni, ai = _fundir(INST, COLUNAS_INST, inst, lambda r: _norm(r.get("entidade")))
    if ciclos:   # sem ciclos lidos (tabela ilegível) não se retira nada
        avisos += _retirar_ausentes(DADOS, COLUNAS_CICLO, ciclos, chave)
    avisos += _limpar_duplicados(DADOS, COLUNAS_CICLO)

    if verbose:
        nomes = [os.path.basename(c) for c in caminhos]
        print("Importação de tabela (Excel/CSV):")
        print("  - fontes lidas ............ %d (%s)" % (len(caminhos), ", ".join(nomes)))
        print("  - ciclos lidos ............ %d" % len(ciclos))
        print("  - ciclos novos/atualizados  %d / %d" % (nc, ac))
        if ni or ai:
            print("  - institucional novos/atu.  %d / %d" % (ni, ai))
        for linha in avisos:
            print(linha)
    return nc + ac + ni + ai


if __name__ == "__main__":
    importar()
