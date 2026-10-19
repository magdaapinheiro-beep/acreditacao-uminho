#!/usr/bin/env python3
"""Gera index.html a partir de dados/acreditacao_uminho.xlsx.

Uso (na raiz do repositório):
    python3 scripts/gerar_pagina.py

Lê as folhas «Ciclos», «Institucional» e «Configuração», converte-as no
formato que a página espera e injeta-as no modelo modelo/pagina.html,
juntamente com o logótipo (assets/logo-uminho.png, embutido em base64 para
que a página funcione como um ficheiro único).
"""
import base64
import datetime as dt
import json
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("Falta o openpyxl. Instale com:  pip install -r requirements.txt")

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ / "scripts"))
import colunas as C  # noqa: E402

EXCEL = RAIZ / "dados" / "acreditacao_uminho.xlsx"
MODELO = RAIZ / "modelo" / "pagina.html"
LOGO = RAIZ / "assets" / "logo-uminho.png"
SAIDA = RAIZ / "index.html"


def valor(v, tipo):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if tipo == "data":
        if isinstance(v, dt.datetime):
            return v.date().isoformat()
        if isinstance(v, dt.date):
            return v.isoformat()
        s = str(v).strip()
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
            try:
                return dt.datetime.strptime(s, fmt).date().isoformat()
            except ValueError:
                pass
        raise ValueError(f"data inválida: {s!r}")
    if tipo == "inteiro":
        return int(float(v))
    return str(v).strip()


def ler_folha(wb, nome, colunas):
    ws = wb[nome]
    cab = [c.value for c in ws[1]]
    esperado = [h for h, *_ in colunas]
    em_falta = [h for h in esperado if h not in cab]
    if em_falta:
        sys.exit(f"Folha «{nome}»: faltam as colunas {em_falta}. Não altere os cabeçalhos.")
    idx = {h: cab.index(h) for h in esperado}
    linhas = []
    for n, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if all(v is None or str(v).strip() == "" for v in row):
            continue
        reg = {}
        for h, campo, tipo, _ in colunas:
            try:
                v = valor(row[idx[h]], tipo)
            except ValueError as e:
                sys.exit(f"Folha «{nome}», linha {n}, coluna «{h}»: {e}")
            if v is not None:
                reg[campo] = v
            elif campo in C.SEMPRE:
                reg[campo] = ""
        linhas.append(reg)
    return linhas


def main():
    wb = load_workbook(EXCEL, data_only=True)
    ciclos = ler_folha(wb, "Ciclos", C.CICLOS)
    inst = ler_folha(wb, "Institucional", C.INSTITUCIONAL)
    data_cfg = valor(wb["Configuração"]["B2"].value, "data") if "Configuração" in wb.sheetnames else None
    atualizado = data_cfg or dt.date.today().isoformat()
    gerado = dt.datetime.now().strftime("%Y-%m-%d %H:%M")

    dados = (
        f'const META = {{ atualizado: "{atualizado}", geradoEm: "{gerado}" }};\n\n'
        f"const INSTITUCIONAL = {json.dumps(inst, ensure_ascii=False, indent=1)};\n\n"
        f"const CICLOS = {json.dumps(ciclos, ensure_ascii=False, indent=1)};\n\n"
    )
    html = MODELO.read_text(encoding="utf-8")
    if "/*__DADOS__*/" not in html or "__LOGO__" not in html:
        sys.exit("O modelo não tem os marcadores /*__DADOS__*/ e __LOGO__.")
    logo = "data:image/png;base64," + base64.b64encode(LOGO.read_bytes()).decode()
    html = html.replace("/*__DADOS__*/\n\n", dados).replace("__LOGO__", logo)
    SAIDA.write_text(html, encoding="utf-8")
    print(f"index.html gerado: {len(ciclos)} ciclos, {len(inst)} registo(s) institucional(is), atualizado a {atualizado}.")


if __name__ == "__main__":
    main()
